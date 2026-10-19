"""
정답지(라벨링 시트)와 확인용 2x2 이미지를 만드는 스크립트.

하는 일:
  1) 날짜마다 매 정시(00:00~23:00)에 가장 가까운 사진(5분 이내)을 카메라별로 고른다
  2) 카메라 4대 사진을 2x2로 합친 확인용 이미지를 저장한다
  3) 정답 엑셀을 만든다 (조명: 켜짐/꺼짐/모름, 사람: 있음/없음/모름 을 드롭다운으로 입력)

폴더 구조: <root>/<room>/<cam1~4>/<YYYY-MM-DD>/<HH-MM-SS>.jpg

사용 예:
  python labels/make_label_sheet.py /Volumes/T9/data 301

결과:
  labels/정답지_301.xlsx
  labels/images/301_<날짜>_<HH>00.jpg
"""
import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image, ImageDraw, ImageFont

CAMS = ["cam1", "cam2", "cam3", "cam4"]
TILE_W, TILE_H = 960, 540  # 2x2로 합치면 1920x1080
NEAR_MIN = 5  # 정시에서 이 시간(분) 안의 사진만 사용


def to_min(stem):
    h, m, s = map(int, stem.split("-"))
    return h * 60 + m + s / 60


def nearest(photos, target):
    """{분: 경로}에서 target(분)에 가장 가까운 사진. 5분 넘게 떨어지면 None"""
    if not photos:
        return None
    k = min(photos, key=lambda t: abs(t - target))
    return photos[k] if abs(k - target) <= NEAR_MIN else None


def load_font(size):
    for f in ["/System/Library/Fonts/Supplemental/Arial Bold.ttf",
              "/System/Library/Fonts/Helvetica.ttc",
              "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"]:
        try:
            return ImageFont.truetype(f, size)
        except OSError:
            continue
    return ImageFont.load_default()


def make_grid(paths, title, out_path, font):
    grid = Image.new("RGB", (TILE_W * 2, TILE_H * 2), (40, 40, 40))
    draw = ImageDraw.Draw(grid)
    for i, cam in enumerate(CAMS):
        x, y = (i % 2) * TILE_W, (i // 2) * TILE_H
        p = paths.get(cam)
        if p:
            grid.paste(Image.open(p).convert("RGB").resize((TILE_W, TILE_H)), (x, y))
            label = f"{cam}  {p.stem.replace('-', ':')}"
        else:
            label = f"{cam}  (no image)"
        draw.rectangle([x, y, x + 330, y + 44], fill=(0, 0, 0))
        draw.text((x + 10, y + 8), label, fill=(255, 255, 0), font=font)
    # 날짜·시각 제목은 아래쪽 가운데 (카메라 이름 표시와 겹치지 않게)
    draw.rectangle([TILE_W - 230, TILE_H * 2 - 52, TILE_W + 230, TILE_H * 2], fill=(0, 0, 0))
    draw.text((TILE_W - 215, TILE_H * 2 - 44), title, fill=(255, 255, 255), font=font)
    grid.save(out_path, quality=85)


def main():
    ap = argparse.ArgumentParser(description="정답지 + 확인용 2x2 이미지 만들기")
    ap.add_argument("root", help="예: /Volumes/T9/data")
    ap.add_argument("room", help="예: 301")
    ap.add_argument("-o", "--out", default=str(Path(__file__).resolve().parent))
    args = ap.parse_args()

    out = Path(args.out)
    img_dir = out / "images"
    img_dir.mkdir(parents=True, exist_ok=True)
    font = load_font(28)
    room_dir = Path(args.root) / args.room

    dates = sorted({d.name for cam in CAMS for d in (room_dir / cam).glob("20*-*-*") if d.is_dir()})
    rows = []
    for date in dates:
        photos = {}
        for cam in CAMS:
            day = room_dir / cam / date
            photos[cam] = {to_min(p.stem): p for p in day.glob("*.jpg")
                           if not p.name.startswith("._")} if day.is_dir() else {}
        for h in range(24):
            paths = {cam: nearest(photos[cam], h * 60) for cam in CAMS}
            paths = {c: p for c, p in paths.items() if p}
            if not paths:  # 그 시각 사진이 하나도 없으면 줄을 만들지 않음
                continue
            name = f"{args.room}_{date}_{h:02d}00.jpg"
            make_grid(paths, f"{args.room}  {date}  {h:02d}:00", img_dir / name, font)
            rows.append((date, f"{h:02d}:00", len(paths), name))
        print(f"{date}: {sum(r[0] == date for r in rows)}개 시각")

    # 엑셀 정답지
    wb = Workbook()
    ws = wb.active
    ws.title = "정답"
    head = ["날짜", "시각", "조명", "사람", "메모", "사진", "카메라 수"]
    ws.append(head)
    for date, hhmm, n, name in rows:
        ws.append([date, hhmm, None, None, None, name, n])
        c = ws.cell(row=ws.max_row, column=6)
        c.hyperlink = f"images/{name}"  # 클릭하면 사진이 열림 (엑셀 파일과 같은 폴더 기준)
        c.font = Font(color="0563C1", underline="single")

    last = ws.max_row
    dv_light = DataValidation(type="list", formula1='"켜짐,꺼짐,모름"', allow_blank=True)
    dv_people = DataValidation(type="list", formula1='"있음,없음,모름"', allow_blank=True)
    ws.add_data_validation(dv_light)
    ws.add_data_validation(dv_people)
    dv_light.add(f"C2:C{last}")
    dv_people.add(f"D2:D{last}")

    fill = PatternFill("solid", fgColor="FFF2CC")  # 입력할 칸 표시
    for r in range(2, last + 1):
        ws.cell(row=r, column=3).fill = fill
        ws.cell(row=r, column=4).fill = fill
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center")
    for col, w in zip("ABCDEFG", [12, 8, 8, 8, 30, 28, 10]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"

    # 입력 기준 안내 시트
    guide = wb.create_sheet("기준")
    for line in [
        ["조명", "강의실 안쪽(바닥·책상)이 고르게 밝으면 '켜짐', 어둡고 칙칙하면 '꺼짐'"],
        ["", "cam1 오른쪽처럼 복도 불빛만 밝은 부분은 판단 근거로 쓰지 않음"],
        ["", "일부만 켜진 것 같거나 애매하면 '모름' + 메모"],
        ["사람", "카메라 4대 중 한 대라도 사람이 보이면 '있음'"],
        ["", "모니터 속 사람, 의자에 걸린 옷, 인형은 사람이 아님"],
        ["", "어두워서 확실하지 않으면 '모름' + 메모"],
    ]:
        guide.append(line)
    guide.column_dimensions["A"].width = 8
    guide.column_dimensions["B"].width = 70

    xlsx = out / f"정답지_{args.room}.xlsx"
    wb.save(xlsx)
    print(f"\n총 {len(rows)}개 시각\n정답지: {xlsx}\n사진: {img_dir}/")


if __name__ == "__main__":
    main()
