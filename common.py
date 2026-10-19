"""
여러 스크립트가 같이 쓰는 함수 모음.

폴더 구조: <root>/<room>/<cam1~4>/<YYYY-MM-DD>/<HH-MM-SS>.jpg
"""
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

CAMS = ["cam1", "cam2", "cam3", "cam4"]
NEAR_MIN = 5   # 정시에서 이 시간(분) 안의 사진만 사용 (정답지와 같은 기준)
GRAY_SAT = 12  # 채도가 이보다 낮으면 흑백(적외선) 사진으로 봄


def to_min(stem):
    """'14-20-03' -> 하루 중 몇 분째인지"""
    h, m, s = map(int, stem.split("-"))
    return h * 60 + m + s / 60


def list_photos(day_dir):
    """{분: 경로}. '._'로 시작하는 파일은 맥이 외장 디스크에 만드는 숨김 파일이라 제외"""
    day_dir = Path(day_dir)
    if not day_dir.is_dir():
        return {}
    return {to_min(p.stem): p for p in day_dir.glob("*.jpg") if not p.name.startswith("._")}


def hourly_photos(root, room, cams=CAMS, dates=None):
    """매 정시에 가장 가까운(5분 이내) 사진을 카메라별로 돌려준다.
    -> [(날짜, 'HH:00', cam, 경로), ...]  정답지의 (날짜, 시각)과 그대로 맞춰짐"""
    room_dir = Path(root) / room
    all_dates = sorted({d.name for c in cams for d in (room_dir / c).glob("20*-*-*") if d.is_dir()})
    out = []
    for date in all_dates:
        if dates and date not in dates:
            continue
        for cam in cams:
            photos = list_photos(room_dir / cam / date)
            if not photos:
                continue
            for h in range(24):
                k = min(photos, key=lambda t: abs(t - h * 60))
                if abs(k - h * 60) <= NEAR_MIN:
                    out.append((date, f"{h:02d}:00", cam, photos[k]))
    return out


def parse_exclude(items):
    """['cam1:0.6,0.3,1,1', ...] -> {'cam1': [(0.6, 0.3, 1.0, 1.0)]}
    사진 크기에 대한 비율(0~1) 좌표: 왼쪽 위 x, y, 오른쪽 아래 x, y"""
    rects = {}
    for it in items or []:
        cam, nums = it.split(":")
        x1, y1, x2, y2 = map(float, nums.split(","))
        rects.setdefault(cam, []).append((x1, y1, x2, y2))
    return rects


def keep_mask(h, w, rects):
    """제외 영역은 False, 나머지는 True인 마스크"""
    m = np.ones((h, w), dtype=bool)
    for x1, y1, x2, y2 in rects:
        m[int(y1 * h):int(y2 * h), int(x1 * w):int(x2 * w)] = False
    return m


def is_grayscale(img):
    sat = np.asarray(img.convert("RGB").convert("HSV"), dtype=np.float32)[:, :, 1]
    return float(sat.mean()) < GRAY_SAT


def read_truth(xlsx, column):
    """정답지 -> {(날짜, 'HH:00'): 값}. '모름'과 빈칸은 뺀다.
    column: '조명' 또는 '사람'"""
    ws = load_workbook(xlsx, data_only=True)["정답"]
    head = [c.value for c in ws[1]]
    i_d, i_t, i_v = head.index("날짜"), head.index("시각"), head.index(column)
    truth = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        d, t, v = row[i_d], row[i_t], row[i_v]
        if d and t and v and v != "모름":
            truth[(str(d)[:10], str(t)[:5])] = str(v).strip()
    return truth
