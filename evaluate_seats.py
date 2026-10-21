"""
person_judge.py 의 좌석별 판정을 좌석 정답지로 채점한다.

정답: labels/좌석정답지_<room>.xlsx  (좌석마다 점유/빈자리/모름)
예측: results/.../seats_<room>.csv  (person_judge.py 결과, 좌석마다 점유/빈자리/판정불가)

채점 규칙:
  - 정답이 '모름'이거나 빈칸인 좌석은 뺀다 (서 있는 사람, 가려서 안 보이는 좌석)
  - 예측이 '판정불가'인 좌석은 --unsure 로 정한다
      skip  : 채점에서 빼고 개수만 보여줌 (기본)
      empty : 빈자리로 봄 (VLM 없이 돌렸을 때 '애매하면 안 잡은 것'으로 보는 경우)
      occ   : 점유로 봄
  - 빈자리가 대부분이라 '정확도'만 보면 아무것도 안 잡아도 높게 나온다.
    그래서 점유를 기준으로 아래 숫자를 같이 본다.
      놓친 점유 (FN) : 정답 점유인데 빈자리로 판정
      헛점유   (FP) : 정답 빈자리인데 점유로 판정  -> 옷, 의자 등받이 같은 오판
      재현율 = 맞힌 점유 / 정답 점유     (사람을 얼마나 놓치지 않는지)
      정밀도 = 맞힌 점유 / 점유라고 한 것 (점유라고 한 게 얼마나 진짜인지)
  - 사용률 통계에 바로 쓰이는 '시각별 점유 좌석 수'가 정답과 몇 석 차이 나는지도 본다.

사용 예:
  python evaluate_seats.py labels/좌석정답지_301.xlsx results/person_301/seats_301.csv
  # 여러 결과를 한 번에 비교 (예: 고치기 전/후, 모델별)
  python evaluate_seats.py labels/좌석정답지_301.xlsx results/person_301_26x_novlm/seats_301.csv results/person_301/seats_301.csv
  python evaluate_seats.py labels/좌석정답지_301.xlsx results/person_301/seats_301.csv --unsure empty

결과:
  화면 출력 + 예측 CSV와 같은 폴더의 eval_seats_<room>.csv (좌석 x 시각마다 정답과 예측)
"""
import argparse
import csv
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

FIXED_COLS = 4  # 날짜, 시각, 묶음, 사진 다음부터 좌석 열


def norm_time(v):
    """엑셀 칸 값이 '15:00', datetime.time, '15:00:00' 어느 것이든 'HH:MM'으로"""
    s = str(v).strip()
    return s[:5] if len(s) >= 5 else s.zfill(5)


def read_seat_truth(xlsx):
    """-> ({(날짜, 'HH:MM'): {좌석: 점유/빈자리/모름}}, {(날짜, 시각): 묶음})"""
    ws = load_workbook(xlsx, data_only=True)["정답"]
    head = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    seat_cols = [(i, h) for i, h in enumerate(head) if i >= FIXED_COLS and h.isdigit()]
    truth, bucket = {}, {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[0] or not row[1]:
            continue
        key = (str(row[0])[:10], norm_time(row[1]))
        truth[key] = {sid: (str(row[i]).strip() if row[i] else "") for i, sid in seat_cols}
        bucket[key] = str(row[2] or "")
    return truth, bucket


def read_pred(path):
    """seats_<room>.csv -> {(날짜, 'HH:MM'): {좌석: (상태, 근거, 잡힌 카메라)}}"""
    pred = defaultdict(dict)
    for r in csv.DictReader(open(path, encoding="utf-8-sig")):
        pred[(r["date"], norm_time(r["time"]))][str(r["seat"])] = (r["상태"], r.get("근거", ""),
                                                                  r.get("잡힌_카메라", ""))
    return pred


def pct(a, b):
    return f"{a / b * 100:5.1f}% ({a}/{b})" if b else "    -"


def score(name, truth, bucket, pred, unsure, out_csv):
    c = Counter()
    by_bucket = defaultdict(Counter)
    count_diff = []      # (키, 정답 점유 수, 예측 점유 수)
    wrong = []           # (키, 좌석, 종류, 근거, 카메라)
    rows = []
    for key in sorted(truth):
        if key not in pred:
            c["예측 없는 시각"] += 1
            continue
        n_true = n_pred = 0
        for sid, t in truth[key].items():
            p, basis, cams = pred[key].get(sid, ("예측없음", "", ""))
            row = {"date": key[0], "time": key[1], "묶음": bucket[key], "seat": sid,
                   "정답": t or "(빈칸)", "예측": p, "근거": basis, "잡힌_카메라": cams, "결과": ""}
            if p == "판정불가":
                c["예측 판정불가"] += 1
                p = {"empty": "빈자리", "occ": "점유"}.get(unsure, p)
            if t not in ("점유", "빈자리"):
                c["정답 모름/빈칸 (제외)"] += 1
                row["결과"] = "제외"
                rows.append(row)
                continue
            if p not in ("점유", "빈자리"):
                c["판정불가로 제외"] += 1
                row["결과"] = "제외(판정불가)"
                rows.append(row)
                continue
            kind = {("점유", "점유"): "TP", ("빈자리", "빈자리"): "TN",
                    ("점유", "빈자리"): "FN", ("빈자리", "점유"): "FP"}[(t, p)]
            c[kind] += 1
            by_bucket[bucket[key]][kind] += 1
            n_true += t == "점유"
            n_pred += p == "점유"
            row["결과"] = {"TP": "맞음(점유)", "TN": "맞음(빈자리)", "FN": "놓친 점유", "FP": "헛점유"}[kind]
            if kind in ("FN", "FP"):
                wrong.append((key, sid, row["결과"], basis, cams))
            rows.append(row)
        count_diff.append((key, n_true, n_pred))

    tp, tn, fn, fp = c["TP"], c["TN"], c["FN"], c["FP"]
    n = tp + tn + fn + fp
    print(f"\n==================== {name}")
    print(f"채점한 좌석 칸 {n}개 (정답 점유 {tp + fn}, 정답 빈자리 {tn + fp})")
    for k in ("정답 모름/빈칸 (제외)", "예측 판정불가", "판정불가로 제외", "예측 없는 시각"):
        if c[k]:
            print(f"  {k}: {c[k]}")
    print(f"  정확도       {pct(tp + tn, n)}   <- 전부 빈자리라고만 해도 {pct(tn + fp, n).strip()}")
    print(f"  재현율(점유) {pct(tp, tp + fn)}   놓친 점유 {fn}")
    print(f"  정밀도(점유) {pct(tp, tp + fp)}   헛점유 {fp}")
    if tp:
        p_, r_ = tp / (tp + fp), tp / (tp + fn)
        print(f"  F1           {2 * p_ * r_ / (p_ + r_) * 100:5.1f}%")

    print("\n  묶음별                정확도             재현율            헛점유")
    for b in sorted(by_bucket):
        k = by_bucket[b]
        print(f"   {b:8s}  {pct(k['TP'] + k['TN'], sum(k.values())):>18s}  "
              f"{pct(k['TP'], k['TP'] + k['FN']):>16s}  {k['FP']:>4d}")

    if count_diff:
        diffs = [p - t for _, t, p in count_diff]
        print(f"\n  시각별 점유 좌석 수: 평균 {sum(abs(d) for d in diffs) / len(diffs):.2f}석 차이, "
              f"정확히 맞은 시각 {sum(d == 0 for d in diffs)}/{len(diffs)}")
        for (d, t), nt, np_ in count_diff:
            if nt != np_:
                print(f"    {d} {t} {bucket[(d, t)]:6s} 정답 {nt:2d}석 / 예측 {np_:2d}석 ({np_ - nt:+d})")

    if wrong:
        print("\n  틀린 좌석:")
        for (d, t), sid, kind, basis, cams in wrong:
            extra = f" [{basis} {cams}]" if basis else ""
            print(f"    {d} {t} {sid:>2s}번 {kind}{extra}")

    with open(out_csv, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()) if rows else ["date"])
        w.writeheader()
        w.writerows(rows)
    print(f"\n  좌석별 결과 CSV: {out_csv}")


def main():
    ap = argparse.ArgumentParser(description="좌석 점유 판정 채점")
    ap.add_argument("truth", help="좌석 정답지, 예: labels/좌석정답지_301.xlsx")
    ap.add_argument("pred", nargs="+", help="person_judge.py 의 seats_<room>.csv (여러 개 가능)")
    ap.add_argument("--unsure", choices=["skip", "empty", "occ"], default="skip",
                    help="예측 '판정불가'를 어떻게 볼지 (기본 skip=채점 제외)")
    args = ap.parse_args()

    truth, bucket = read_seat_truth(args.truth)
    room = Path(args.truth).stem.split("_")[-1]
    print(f"정답 {len(truth)}개 시각, 좌석 {len(next(iter(truth.values()), {}))}개")
    for p in args.pred:
        pred = read_pred(p)
        missing = [k for k in truth if k not in pred]
        if missing:
            print(f"\n주의: {p} 에 없는 정답 시각 {len(missing)}개: "
                  + ", ".join(f"{d} {t}" for d, t in missing[:5]) + (" ..." if len(missing) > 5 else ""))
        score(p, truth, bucket, pred, args.unsure, Path(p).parent / f"eval_seats_{room}.csv")


if __name__ == "__main__":
    main()