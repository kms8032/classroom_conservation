"""
[광각 인원수] 광각 카메라(cam5) 사진 한 장으로 강의실 전체 인원을 센다.

광각 사진(3632x1632)의 특징과 대응:
  - 위쪽 약 1/3이 천장 -> --crop-top 비율만큼 잘라내고 탐지 (사람이 있을 수 없는 곳에 픽셀을 쓰지 않음)
  - 앞줄 사람이 아주 작음 -> 사진을 가로로 N조각(--tiles) 나눠서 각각 탐지한 뒤 합침 (타일링)
       조각끼리 조금 겹치게 자르고, 겹친 곳에서 두 번 잡힌 사람은 하나로 합친다.
  - 신뢰도: conf >= 0.5 는 '확정', 0.25~0.5 는 '애매' 로 따로 센다.

세 단계로 쓴다:
  1) 탐지:  python count_wide.py detect /Volumes/T9/data 301
            python count_wide.py detect /Volumes/T9/data 301 --tiles 3 -o results/wide_301_t3 --draw
  2) 정답지: python count_wide.py sheet /Volumes/T9/data 301
            -> labels/인원정답지_301.xlsx 의 '인원' 칸에 사진을 보고 실제 인원을 적는다 (탐지 결과는 보지 말 것)
  3) 채점:  python count_wide.py score labels/인원정답지_301.xlsx results/wide_301/counts_301.csv results/wide_301_t3/counts_301.csv

결과 (detect, 기본 results/wide_<room>/):
  counts_<room>.csv  사진마다 한 줄: 확정 인원, 애매 박스 수, 적외선 여부
  boxes_<room>.csv   박스마다 한 줄: 신뢰도, 좌표 (나중에 VLM 확인용)
  images/            --draw 를 주면 박스를 그린 확인용 사진 (확정=초록, 애매=주황)
"""
import argparse
import csv
import math
from pathlib import Path

import cv2
import numpy as np

from common import GRAY_SAT, list_photos


# ---------------------------------------------------------------- 사진 고르기

def wide_photos(root, room, cam, dates=None):
    """-> [(날짜, 'HH:MM', 경로), ...] 시간 순서"""
    out = []
    for day in sorted((Path(root) / room / cam).glob("20*-*-*")):
        if dates and day.name not in dates:
            continue
        for minute, p in sorted(list_photos(day).items()):
            out.append((day.name, f"{int(minute // 60):02d}:{int(minute % 60):02d}", p))
    return out


def is_infrared(img):
    small = cv2.resize(img, (320, 144))
    return float(cv2.cvtColor(small, cv2.COLOR_BGR2HSV)[:, :, 1].mean()) < GRAY_SAT


# ---------------------------------------------------------------- 탐지 (타일링)

def overlap_small(a, b):
    """겹친 넓이 / 작은 박스 넓이. 조각 경계에서 잘린 사람(작은 박스)이 큰 박스와 겹치면 1에 가까움"""
    x1, y1, x2, y2 = max(a[0], b[0]), max(a[1], b[1]), min(a[2], b[2]), min(a[3], b[3])
    inter = max(0, x2 - x1) * max(0, y2 - y1)
    small = min((a[2] - a[0]) * (a[3] - a[1]), (b[2] - b[0]) * (b[3] - b[1]))
    return inter / small if small else 0.0


def merge(dets, thr=0.6):
    """신뢰도 높은 순으로, 이미 고른 박스와 많이 겹치면 같은 사람으로 보고 하나로 합침.
    조각 경계에서 잘린 박스와 온전한 박스가 합쳐지도록 두 박스를 모두 감싸는 박스로 키운다."""
    kept = []
    for d in sorted(dets, key=lambda d: -d["conf"]):
        same = next((k for k in kept if overlap_small(d["box"], k["box"]) >= thr), None)
        if same is None:
            kept.append(dict(d))
        else:
            a, b = same["box"], d["box"]
            same["box"] = (min(a[0], b[0]), min(a[1], b[1]), max(a[2], b[2]), max(a[3], b[3]))
    return kept


def detect(model, img, imgsz, conf_low, crop_top, tiles, overlap, device):
    h, w = img.shape[:2]
    y0 = int(h * crop_top)
    region = img[y0:]
    if tiles <= 1:
        spans = [(0, w)]
    else:
        tw = math.ceil(w / tiles * (1 + overlap))       # 조각 너비 (겹치는 만큼 넓게)
        step = (w - tw) / (tiles - 1)
        spans = [(int(i * step), int(i * step) + tw) for i in range(tiles)]
    dets = []
    for x0, x1 in spans:
        r = model(region[:, x0:x1], imgsz=imgsz, conf=conf_low, classes=[0], verbose=False,
                  device=device or None)[0]
        if r.boxes is None:
            continue
        for b in r.boxes:
            bx1, by1, bx2, by2 = map(float, b.xyxy[0].tolist())
            dets.append({"box": (int(bx1 + x0), int(by1 + y0), int(bx2 + x0), int(by2 + y0)),
                         "conf": round(float(b.conf.item()), 3)})
    return merge(dets) if tiles > 1 else dets


def draw(img, dets, conf_high, crop_top, path):
    out = img.copy()
    y0 = int(img.shape[0] * crop_top)
    cv2.line(out, (0, y0), (img.shape[1], y0), (200, 200, 200), 2)  # 이 선 위는 탐지하지 않음
    for d in dets:
        color = (0, 200, 0) if d["conf"] >= conf_high else (0, 165, 255)
        x1, y1, x2, y2 = d["box"]
        cv2.rectangle(out, (x1, y1), (x2, y2), color, 3)
        cv2.putText(out, f'{d["conf"]:.2f}', (x1, max(20, y1 - 8)), cv2.FONT_HERSHEY_SIMPLEX, 1.0, color, 2)
    n = sum(d["conf"] >= conf_high for d in dets)
    cv2.putText(out, f"{n} (+{len(dets) - n})", (30, 80), cv2.FONT_HERSHEY_SIMPLEX, 2.5, (0, 0, 255), 5)
    path.parent.mkdir(parents=True, exist_ok=True)
    cv2.imwrite(str(path), out, [cv2.IMWRITE_JPEG_QUALITY, 85])


def cmd_detect(args):
    from ultralytics import YOLO
    dates = {d.strip() for d in args.dates.split(",") if d.strip()} or None
    photos = wide_photos(args.root, args.room, args.cam, dates)
    out = Path(args.out or f"results/wide_{args.room}")
    out.mkdir(parents=True, exist_ok=True)
    print(f"{args.cam} 사진 {len(photos)}장 -> YOLO({args.model}, imgsz {args.imgsz}, "
          f"천장 {args.crop_top:.0%} 자름, 조각 {args.tiles}개)")
    model = YOLO(args.model)
    counts, boxes = [], []
    for i, (date, hhmm, p) in enumerate(photos, 1):
        img = cv2.imread(str(p))
        if img is None:
            print(f"  읽기 실패: {p}")
            continue
        dets = detect(model, img, args.imgsz, args.conf_low, args.crop_top, args.tiles, args.overlap, args.device)
        sure = [d for d in dets if d["conf"] >= args.conf_high]
        ir = int(is_infrared(img))
        counts.append({"date": date, "time": hhmm, "확정": len(sure), "애매": len(dets) - len(sure),
                       "적외선": ir, "file": str(p)})
        for d in dets:
            boxes.append({"date": date, "time": hhmm, "conf": d["conf"], "box": "%d,%d,%d,%d" % d["box"],
                          "확정": int(d["conf"] >= args.conf_high), "file": str(p)})
        if args.draw:
            draw(img, dets, args.conf_high, args.crop_top, out / "images" / f"{date}_{hhmm.replace(':', '')}.jpg")
        print(f"  [{i}/{len(photos)}] {date} {hhmm} {'(적외선) ' if ir else ''}확정 {len(sure)}명, 애매 {len(dets) - len(sure)}")
    for name, rows in (("counts", counts), ("boxes", boxes)):
        path = out / f"{name}_{args.room}.csv"
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.DictWriter(f, fieldnames=list(rows[0].keys()) if rows else ["date"])
            w.writeheader()
            w.writerows(rows)
        print(f"CSV: {path} ({len(rows)}줄)")


# ---------------------------------------------------------------- 정답지

def cmd_sheet(args):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    xlsx = Path(args.out or f"labels/인원정답지_{args.room}.xlsx")
    if xlsx.exists():
        raise SystemExit(f"{xlsx} 가 이미 있음. 적어 둔 정답 보호를 위해 멈춤 (지우거나 이름을 바꾼 뒤 실행)")
    photos = wide_photos(args.root, args.room, args.cam)
    wb = Workbook()
    ws = wb.active
    ws.title = "정답"
    ws.append(["날짜", "시각", "사진", "인원", "상태", "메모"])
    for date, hhmm, p in photos:
        ws.append([date, hhmm, p.name, None, None, None])
        c = ws.cell(ws.max_row, 3)
        c.hyperlink = Path(p).resolve().as_uri()  # 클릭하면 외장하드의 원본 사진이 열림
        c.font = Font(color="0563C1", underline="single")
    last = ws.max_row
    dv_n = DataValidation(type="whole", operator="between", formula1="0", formula2="60", allow_blank=True)
    dv_s = DataValidation(type="list", formula1='"빈 강의실,자습,수업,모름"', allow_blank=True)
    ws.add_data_validation(dv_n)
    ws.add_data_validation(dv_s)
    dv_n.add(f"D2:D{last}")
    dv_s.add(f"E2:E{last}")
    fill = PatternFill("solid", fgColor="FFF2CC")
    for r in range(2, last + 1):
        ws.cell(r, 4).fill = fill
        ws.cell(r, 4).alignment = Alignment(horizontal="center")
    for c in ws[1]:
        c.font = Font(bold=True)
    for col, wdt in zip("ABCDEF", [12, 7, 16, 7, 11, 40]):
        ws.column_dimensions[col].width = wdt
    ws.freeze_panes = "A2"

    g = wb.create_sheet("기준")
    for line in [
        ["원칙", "count_wide.py 탐지 결과(그림)를 보지 말고 원본 사진만 보고 센다"],
        ["인원", "사진에 보이는 사람 전부 (앉은 사람, 서 있는 사람, 강사 포함). 모니터 위로 머리만 보여도 1명"],
        ["", "의자에 걸린 옷, 가방, 인형, 모니터 화면 속 사람은 세지 않는다"],
        ["", "확실하지 않은 사람이 있으면 확실한 수만 적고 메모에 '+1 애매' 처럼 남긴다"],
        ["상태", "빈 강의실: 0명 / 자습: 강사 없이 각자 공부 / 수업: 강사가 있거나 강의 화면을 다 같이 봄"],
        ["", "판단이 어려우면 '모름'. 상태는 비워 둬도 됨 (인원이 먼저)"],
    ]:
        g.append(line)
    g.column_dimensions["A"].width = 8
    g.column_dimensions["B"].width = 90
    xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx)
    print(f"정답지: {xlsx} ({len(photos)}장)")


# ---------------------------------------------------------------- 채점

def read_counts_truth(xlsx):
    from openpyxl import load_workbook
    ws = load_workbook(xlsx, data_only=True)["정답"]
    truth = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        d, t, n = row[0], row[1], row[3]
        if d and t is not None and n is not None and str(n).strip() != "":
            truth[(str(d)[:10], str(t)[:5].zfill(5))] = int(n)
    return truth


def cmd_score(args):
    truth = read_counts_truth(args.truth)
    print(f"정답 {len(truth)}장 (실제 인원 합 {sum(truth.values())}명, 0명인 사진 {sum(v == 0 for v in truth.values())}장)")
    for path in args.counts:
        pred = {(r["date"], r["time"]): r for r in csv.DictReader(open(path, encoding="utf-8-sig"))}
        keys = [k for k in sorted(truth) if k in pred]
        print(f"\n==================== {path}  (채점 {len(keys)}장)")
        for label, f in (("확정만", lambda r: int(r["확정"])), ("확정+애매", lambda r: int(r["확정"]) + int(r["애매"]))):
            diff = [f(pred[k]) - truth[k] for k in keys]
            if not diff:
                continue
            print(f"  [{label}] 정확히 맞음 {sum(d == 0 for d in diff)}/{len(diff)}장, "
                  f"평균 {sum(abs(d) for d in diff) / len(diff):.2f}명 차이, "
                  f"적게 셈 {-sum(d for d in diff if d < 0)}명, 많이 셈 {sum(d for d in diff if d > 0)}명")
        print("  틀린 사진 (확정만 기준):")
        for k in keys:
            p = int(pred[k]["확정"])
            if p != truth[k]:
                print(f"    {k[0]} {k[1]} 정답 {truth[k]:2d}명 / 탐지 {p:2d}명 ({p - truth[k]:+d})"
                      f"  애매 {pred[k]['애매']}{'  적외선' if pred[k]['적외선'] == '1' else ''}")


def main():
    ap = argparse.ArgumentParser(description="광각 카메라 인원수 세기")
    sub = ap.add_subparsers(dest="cmd", required=True)

    d = sub.add_parser("detect", help="YOLO로 인원 세기")
    d.add_argument("root", help="예: /Volumes/T9/data")
    d.add_argument("room", help="예: 301")
    d.add_argument("--cam", default="cam5")
    d.add_argument("--dates", default="", help="쉼표로 구분. 비우면 전체")
    d.add_argument("--model", default="yolo26x-pose.pt")
    d.add_argument("--imgsz", type=int, default=1280)
    d.add_argument("--crop-top", type=float, default=0.35, help="위에서 이 비율만큼 잘라냄 (천장). 0이면 안 자름")
    d.add_argument("--tiles", type=int, default=1, help="가로로 나눌 조각 수 (1=나누지 않음)")
    d.add_argument("--overlap", type=float, default=0.15, help="조각끼리 겹치는 비율")
    d.add_argument("--conf-high", type=float, default=0.5)
    d.add_argument("--conf-low", type=float, default=0.25)
    d.add_argument("--device", default="", help="mps, cpu, 0. 비우면 자동")
    d.add_argument("--draw", action="store_true", help="확인용 그림 저장")
    d.add_argument("-o", "--out", default="")

    s = sub.add_parser("sheet", help="인원 정답지 만들기")
    s.add_argument("root")
    s.add_argument("room")
    s.add_argument("--cam", default="cam5")
    s.add_argument("-o", "--out", default="")

    c = sub.add_parser("score", help="정답지로 채점 (counts CSV 여러 개 가능)")
    c.add_argument("truth")
    c.add_argument("counts", nargs="+")

    args = ap.parse_args()
    {"detect": cmd_detect, "sheet": cmd_sheet, "score": cmd_score}[args.cmd](args)


if __name__ == "__main__":
    main()